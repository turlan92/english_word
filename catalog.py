"""Catalog migration and validated, append-only imports. No Telegram dependency."""
import csv
import io
import sqlite3
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile, BadZipFile

HEADERS = ['книга', 'глава', 'слово', 'перевод', 'пример', 'перевод примера']

@contextmanager
def database(path):
    db = sqlite3.connect(path, timeout=15)
    db.row_factory = sqlite3.Row
    try:
        with db:
            yield db
    finally:
        db.close()

def backup(path):
    path = Path(path)
    if not path.exists():
        return
    folder = path.parent / 'backups'
    folder.mkdir(exist_ok=True)
    target = folder / (path.stem + '_' + datetime.now().strftime('%Y%m%d_%H%M%S_%f') + '.sqlite3')
    with database(path) as src, database(target) as dst:
        src.backup(dst)

def migrate(path):
    with database(path) as db:
        exists = db.execute("SELECT 1 FROM sqlite_master WHERE name='catalog'").fetchone()
    if exists:
        return
    backup(path)
    with database(path) as db:
        db.execute('''CREATE TABLE collections (
            section TEXT PRIMARY KEY, book TEXT NOT NULL, chapter TEXT NOT NULL,
            UNIQUE(book,chapter))''')
        db.execute('''CREATE TABLE catalog (
            id INTEGER PRIMARY KEY, section TEXT NOT NULL, verb TEXT NOT NULL,
            meaning_ru TEXT NOT NULL, definition_en TEXT, source_page INTEGER,
            sort_order INTEGER NOT NULL, needs_review INTEGER NOT NULL DEFAULT 0,
            example TEXT NOT NULL DEFAULT '', example_ru TEXT NOT NULL DEFAULT '',
            UNIQUE(section,sort_order))''')
        db.executemany('INSERT INTO collections VALUES(?,?,?)', [
            ('basic', 'Васильева — 100 Basic Phrasal Verbs', 'Основные 100'),
            ('extended', 'Васильева — 100 Basic Phrasal Verbs', 'Расширенный словарь')])
        db.execute('''INSERT INTO catalog(id,section,verb,meaning_ru,definition_en,source_page,sort_order,needs_review)
            SELECT id,section,verb,meaning_ru,definition_en,source_page,sort_order,needs_review FROM entries''')

def parse_import(data, filename):
    if len(data) > 2_000_000:
        raise ValueError('Файл должен быть не больше 2 МБ.')
    if filename.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        try:
            with ZipFile(io.BytesIO(data)) as z:
                if sum(x.file_size for x in z.infolist()) > 20_000_000:
                    raise ValueError('Слишком большой распакованный Excel-файл.')
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            try:
                rows = []
                for row in wb.worksheets[0].iter_rows():
                    if len(rows) > 1000:
                        raise ValueError('Не больше 1000 строк за загрузку.')
                    if any(c.data_type == 'f' for c in row):
                        raise ValueError('Формулы не допускаются. Вставьте значения как текст.')
                    rows.append([c.value for c in row])
            finally:
                wb.close()
        except (BadZipFile, KeyError, OSError) as exc:
            raise ValueError('Не удалось прочитать XLSX.') from exc
    elif filename.lower().endswith('.csv'):
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ValueError('Сохраните CSV в кодировке UTF-8.') from exc
        delimiter = ';' if ';' in text.split('\n')[0] else ','
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    else:
        raise ValueError('Поддерживаются .xlsx и .csv.')
    rows = [r for r in rows if any(str(v or '').strip() for v in r)]
    if not rows or [str(x or '').strip().lower() for x in rows[0][:6]] != HEADERS:
        raise ValueError('Первая строка: ' + '; '.join(HEADERS))
    if len(rows) > 1001:
        raise ValueError('Не больше 1000 слов за загрузку.')
    result = []
    for i, raw in enumerate(rows[1:], 2):
        vals = [str(v).strip() if v is not None else '' for v in raw]
        vals += [''] * max(0, 6 - len(vals))
        if any(vals[6:]):
            raise ValueError(f'Строка {i}: лишние столбцы.')
        vals = vals[:6]
        if not all(vals[:4]):
            raise ValueError(f'Строка {i}: заполните книгу, главу, слово и перевод.')
        limits = [120, 120, 150, 350, 600, 600]
        if any(len(v) > limit for v, limit in zip(vals, limits)):
            raise ValueError(f'Строка {i}: слишком длинное поле (лимиты: {limits}).')
        result.append(vals)
    if not result:
        raise ValueError('Нет слов для импорта.')
    return result

def import_rows(path, rows):
    backup(path)
    added = skipped = 0
    with database(path) as db:
        db.execute('BEGIN IMMEDIATE')
        for book, chapter, verb, meaning, example, example_ru in rows:
            section = db.execute('SELECT section FROM collections WHERE book=? AND chapter=?', (book, chapter)).fetchone()
            if section:
                section = section[0]
            else:
                n = db.execute('SELECT COUNT(*) FROM collections').fetchone()[0] + 1
                section = f'set{n}'
                db.execute('INSERT INTO collections VALUES(?,?,?)', (section, book, chapter))
            known = db.execute('SELECT verb,meaning_ru FROM catalog WHERE section=?', (section,)).fetchall()
            norm = lambda s: ' '.join(s.split()).casefold()
            if any(norm(r['verb']) == norm(verb) and norm(r['meaning_ru']) == norm(meaning) for r in known):
                skipped += 1
                continue
            order = db.execute('SELECT COALESCE(MAX(sort_order),0)+1 FROM catalog WHERE section=?', (section,)).fetchone()[0]
            db.execute('''INSERT INTO catalog(section,verb,meaning_ru,sort_order,example,example_ru)
                VALUES(?,?,?,?,?,?)''', (section, verb, meaning, order, example, example_ru))
            added += 1
    return added, skipped
